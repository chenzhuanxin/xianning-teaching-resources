# -*- coding: utf-8 -*-
"""
期末真题试卷 渲染引擎（Word .docx + HTML 双版本）
风格：A4 试卷排版，密封线、题号、分值、答题线、答案页齐全
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

WORD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "交付物", "03_期末真题试卷", "Word_PDF")
HTML_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "交付物", "03_期末真题试卷", "HTML")
os.makedirs(WORD_DIR, exist_ok=True)
os.makedirs(HTML_DIR, exist_ok=True)


# ============================================================
#                    HTML 试卷模板
# ============================================================
PAPER_CSS = """
*{margin:0;padding:0;box-sizing:border-box;}
body{font-family:"SimSun","Microsoft YaHei",serif;background:#e9edf3;color:#000;padding:24px 0;}
.paper{width:21cm;min-height:29.7cm;margin:0 auto 26px;background:#fff;
padding:1.6cm 1.5cm 1.8cm;box-shadow:0 6px 26px rgba(0,0,0,.14);position:relative;}
.paper.ans{background:#fcfcfa;}

/* 密封线 */
.seal{position:absolute;left:.75cm;top:1.6cm;bottom:1.8cm;width:.38cm;
border-right:1px dashed #999;display:flex;align-items:center;justify-content:center;}
.seal span{writing-mode:vertical-rl;font-size:11px;color:#666;letter-spacing:5px;
font-family:"SimSun",serif;white-space:nowrap;}
.seal span b{color:#000;font-weight:700;}

.head{text-align:center;margin-bottom:8px;}
.head .region{font-size:12.5px;color:#444;letter-spacing:2px;margin-bottom:5px;font-family:"SimHei",sans-serif;}
.head h1{font-size:22px;font-weight:800;letter-spacing:2px;font-family:"SimHei",sans-serif;margin-bottom:5px;}
.head .info{font-size:12.5px;color:#333;letter-spacing:.5px;}
.head .rule{border-top:2.5px solid #000;border-bottom:1px solid #000;height:4px;margin:9px 0 6px;}

.scoreline{width:100%;border-collapse:collapse;margin:8px 0 14px;font-size:12px;}
.scoreline td{border:1px solid #333;padding:5px 7px;text-align:center;font-family:"SimHei",sans-serif;}
.scoreline td.lbl{background:#f3f3f3;width:64px;font-weight:700;}
.scoreline td.get{width:74px;font-weight:700;}

.fillinfo{display:flex;justify-content:space-between;font-size:13px;margin:0 0 14px;
border-bottom:1px dashed #999;padding-bottom:11px;}
.fillinfo .u{letter-spacing:1px;}

.part{margin:16px 0 7px;font-size:14.5px;font-weight:800;font-family:"SimHei",sans-serif;
background:#f5f5f5;border-left:5px solid #333;padding:6px 11px;letter-spacing:.5px;}
.tip{font-size:11.8px;color:#555;margin:-3px 0 9px;font-style:italic;}

.q{margin:0 0 9px;font-size:13.6px;line-height:1.95;text-align:justify;}
.q .no{font-weight:800;font-family:"SimHei",sans-serif;}
.q .sc{color:#555;font-size:12.2px;}
.opts{margin:3px 0 3px 24px;font-size:13.4px;line-height:1.92;}
.opts div{margin-bottom:1px;}
.blank{display:inline-block;min-width:56px;border-bottom:1px solid #000;text-align:center;margin:0 3px;}
.bline{border-bottom:1px solid #aaa;height:26px;margin:5px 0 0;}

table.data{width:100%;border-collapse:collapse;margin:7px 0 9px;font-size:13px;}
table.data th,table.data td{border:1px solid #333;padding:6px 8px;text-align:center;}
table.data th{background:#f3f3f3;font-family:"SimHei",sans-serif;font-weight:700;}

.foot{text-align:center;font-size:11.5px;color:#777;margin-top:18px;
border-top:1px dashed #bbb;padding-top:10px;}

/* 答案页 */
.anshead{text-align:center;font-size:17px;font-weight:800;font-family:"SimHei",sans-serif;
margin-bottom:11px;padding-bottom:8px;border-bottom:2.5px solid #000;}
.ans .q{margin-bottom:5px;font-size:13.2px;}
.ans .q .no{color:#000;}
.ans .val{color:#000;font-weight:700;}
.ans .analy{color:#444;font-size:12.6px;}

@media print{
  body{background:#fff;padding:0;}
  .paper{width:auto;min-height:auto;margin:0;box-shadow:none;padding:1.2cm 1.3cm;}
  .paper{page-break-after:always;}
  .paper:last-child{page-break-after:auto;}
}
@media screen and (max-width:900px){
  .paper{width:auto;padding:16px;}
  .seal{display:none;}
}
"""

def _q(num, text, score=None, opts=None, blanks=0, table=None, kind=""):
    """构造一道题"""
    sc = f' <span class="sc">（{score}分）</span>' if score else ""
    kd = f'[{kind}] ' if kind else ""
    h = f'<div class="q"><span class="no">{num}．</span>{kd}{text}{sc}</div>'
    if opts:
        h += '<div class="opts">' + "".join(f"<div>{o}</div>" for o in opts) + "</div>"
    if table:
        h += table
    if blanks:
        h += "".join('<div class="bline"></div>' for _ in range(blanks))
    return h

def _qhtml(q):
    """dict → HTML 题目"""
    if isinstance(q, str):
        return q
    sc = f' <span class="sc">（{q["score"]}分）</span>' if q.get("score") else ""
    kd = f'[{q["kind"]}] ' if q.get("kind") else ""
    h = f'<div class="q"><span class="no">{q["no"]}．</span>{kd}{q["text"]}{sc}</div>'
    if q.get("opts"):
        h += '<div class="opts">' + "".join(f"<div>{o}</div>" for o in q["opts"]) + "</div>"
    for _ in range(q.get("blanks", 0)):
        h += '<div class="bline"></div>'
    return h

def _ahtml(a):
    """答案 dict → HTML"""
    if isinstance(a, str):
        return f'<div class="q">{a}</div>'
    s = f'<div class="q"><span class="no">{a["no"]}．</span>{a.get("kind","")} '
    s += f'<span class="val">{a["val"]}</span></div>'
    if a.get("analy"):
        s += f'<div class="q analy">【解析】{a["analy"]}</div>'
    return s

def render_paper_html(meta, parts, answers=None):
    """
    meta: dict(region, title, subtitle, info, full_score, duration, subject_grade)
    parts: list of dict(name, tip, questions=[dict|str,...])
    answers: list of dict|str
    """
    q_html = ""
    for p in parts:
        q_html += f'<div class="part">{p["name"]}</div>'
        if p.get("tip"):
            q_html += f'<div class="tip">{p["tip"]}</div>'
        q_html += "".join(_qhtml(q) for q in p["questions"])

    body1 = f'''<div class="paper">
<div class="seal"><span>密 封 线 内 不 要 答 题 · <b>{meta["subject_grade"]}</b></span></div>
<div class="head">
  <div class="region">{meta["region"]}</div>
  <h1>{meta["title"]}</h1>
  <div class="info">{meta["info"]}</div>
  <div class="rule"></div>
  <div class="info">满分：{meta["full_score"]}分　　考试时间：{meta["duration"]}分钟</div>
</div>
<table class="scoreline">
  <tr><td class="lbl">题号</td>{ "".join(f"<td>{i}</td>" for i in range(1,len(parts)+2)) }<td class="get">总分</td></tr>
  <tr><td class="lbl">得分</td>{ "".join("<td></td>" for _ in range(len(parts)+1)) }<td></td></tr>
</table>
<div class="fillinfo">
  <div class="u">学校：________________</div>
  <div class="u">班级：____________</div>
  <div class="u">姓名：____________</div>
  <div class="u">考号：____________</div>
</div>
{q_html}
<div class="foot">{meta["region"]} · {meta["subject_grade"]} 期末质量检测试卷 · 第 1 页</div>
</div>'''

    body2 = ""
    if answers:
        body2 = f'''<div class="paper ans">
<div class="anshead">{meta["subject_grade"]} {meta["title"]} · 参考答案与解析</div>
<div class="ans">{"".join(_ahtml(a) for a in answers)}</div>
<div class="foot">参考答案 · 仅供教师评卷与学生自查使用</div>
</div>'''

    return f'''<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>{meta["subject_grade"]} {meta["title"]}</title>
<style>{PAPER_CSS}</style></head><body>
{body1}
{body2}
</body></html>'''


# ============================================================
#                    Word 试卷生成
# ============================================================
def _set_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

def _add_para(doc, text="", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, bold=False,
              name="宋体", space_after=2, space_before=0, indent=None, line=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if line:
        pf.line_spacing = line
    if text:
        r = p.add_run(text)
        _set_font(r, name, size, bold)
    return p

def _add_hr(doc, ch="─", size=9, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = _add_para(doc, ch * 46, align=align, size=size, space_after=4, space_before=4)
    return p

def gen_word_paper(meta, parts, answers, filepath):
    doc = Document()
    # 页面设置 A4
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = Cm(1.7); sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)

    # 页眉（密封线提示）
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run(f"密封线内不要答题　·　{meta['subject_grade']}")
    _set_font(hr, "宋体", 9, False, RGBColor(0x66, 0x66, 0x66))

    # 标题区
    _add_para(doc, meta["region"], WD_ALIGN_PARAGRAPH.CENTER, 11, False, "黑体", 2)
    _add_para(doc, meta["title"], WD_ALIGN_PARAGRAPH.CENTER, 17, True, "黑体", 3)
    _add_para(doc, meta["info"], WD_ALIGN_PARAGRAPH.CENTER, 10.5, False, "宋体", 2)
    _add_para(doc, f"满分：{meta['full_score']}分　　考试时间：{meta['duration']}分钟",
              WD_ALIGN_PARAGRAPH.CENTER, 10.5, False, "宋体", 6)

    # 得分表
    ncol = len(parts) + 2
    tbl = doc.add_table(rows=2, cols=ncol)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["题号"] + [str(i) for i in range(1, len(parts) + 1)] + ["总分"]
    for j, t in enumerate(hdr):
        c = tbl.cell(0, j); c.text = ""
        r = c.paragraphs[0].add_run(t); r.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(r, "黑体", 9.5, True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for j in range(ncol):
        c = tbl.cell(1, j); c.text = ""
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if j == 0:
            r = c.paragraphs[0].add_run("得分"); _set_font(r, "黑体", 9.5, True)
    _add_para(doc, "", space_after=4)

    # 考生信息
    _add_para(doc, "学校：__________________　班级：____________　姓名：____________　考号：____________",
              WD_ALIGN_PARAGRAPH.LEFT, 10.5, False, "宋体", 8)

    # 题目
    for p in parts:
        pp = _add_para(doc, p["name"], WD_ALIGN_PARAGRAPH.LEFT, 11.5, True, "黑体", 3, 8)
        if p.get("tip"):
            _add_para(doc, p["tip"], WD_ALIGN_PARAGRAPH.LEFT, 9, False, "楷体", 5)
        for q in p["questions"]:
            if isinstance(q, dict):
                # 题目文本
                txt = f"{q['no']}．"
                if q.get("kind"):
                    txt += f"[{q['kind']}] "
                txt += q["text"]
                if q.get("score"):
                    txt += f"（{q['score']}分）"
                _add_para(doc, txt, WD_ALIGN_PARAGRAPH.LEFT, 10.5, False, "宋体", 2, 3)
                for o in q.get("opts", []):
                    _add_para(doc, o, WD_ALIGN_PARAGRAPH.LEFT, 10.5, False, "宋体", 1, 0, indent=18)
                for _ in range(q.get("blanks", 0)):
                    _add_para(doc, "＿" * 34, WD_ALIGN_PARAGRAPH.LEFT, 10.5, False, "宋体", 3, 1)
            else:
                _add_para(doc, str(q), WD_ALIGN_PARAGRAPH.LEFT, 10.5, False, "宋体", 3)

    # 答案页
    doc.add_page_break()
    _add_para(doc, f"{meta['subject_grade']} {meta['title']}　参考答案与解析",
              WD_ALIGN_PARAGRAPH.CENTER, 14, True, "黑体", 8)
    for a in answers:
        if isinstance(a, dict):
            _add_para(doc, f"{a['no']}．{a.get('kind','')}　{a['val']}",
                      WD_ALIGN_PARAGRAPH.LEFT, 10.5, False, "宋体", 2)
            if a.get("analy"):
                _add_para(doc, f"【解析】{a['analy']}", WD_ALIGN_PARAGRAPH.LEFT, 9.5,
                          False, "楷体", 4, 0, indent=21)
        else:
            _add_para(doc, str(a), WD_ALIGN_PARAGRAPH.LEFT, 10, False, "宋体", 3)

    doc.save(filepath)
    return filepath
